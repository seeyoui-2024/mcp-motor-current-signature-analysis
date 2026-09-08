"""DOCX report generation for MCSA diagnostics.

Generates a professional Word document report from MCSA analysis results.
Supports English and Chinese (zh) languages.  Mirrors the content of the
HTML diagnostic report in a printable, email-friendly DOCX format.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from mcp_server_mcsa.i18n import Language, get_assessment_text, get_text

# ---------------------------------------------------------------------------
# Severity helpers
# ---------------------------------------------------------------------------

_SEVERITY_COLORS: dict[str, str] = {
    "healthy": "#28a745",
    "incipient": "#ffc107",
    "moderate": "#fd7e14",
    "severe": "#dc3545",
}


def _severity_key(lang: Language) -> dict[str, str]:
    """Return a canonical→label mapping for the four severity levels."""
    return {k: get_text(f"severity.{k}", lang) for k in ("healthy", "incipient", "moderate", "severe")}


def _build_sev_lookup() -> dict[str, str]:
    """Build a reverse lookup: translated label → canonical severity key."""
    lookup: dict[str, str] = {}
    for lang in ("en", "zh"):
        for canonical, label in _severity_key(lang).items():
            lookup[label] = canonical
    return lookup


_SEV_LOOKUP = _build_sev_lookup()


def _severity_color(severity_label: str) -> str:
    canonical = _SEV_LOOKUP.get(severity_label, "")
    return _SEVERITY_COLORS.get(canonical, "#6c757d")


def _find_assessment_key(text: str) -> str:
    """Reverse-map an overall assessment text to its i18n key."""
    upper = text.upper()
    key_map = {
        "CRITICAL": "assessment.critical",
        "WARNING": "assessment.warning",
        "WATCH": "assessment.watch",
        "NORMAL": "assessment.normal",
    }
    for prefix, key in key_map.items():
        if upper.startswith(prefix):
            return key
    return ""


def _get_recommendations(summary: dict, lang: Language) -> list[str]:
    """Build actionable recommendations from the summary severities."""
    recs: list[str] = []
    for sev_key, fault_key in [
        ("brb_severity", "fault_type.broken_rotor_bars"),
        ("eccentricity_severity", "fault_type.eccentricity"),
        ("stator_severity", "fault_type.stator_inter_turn"),
    ]:
        sev_label = summary.get(sev_key, "")
        canonical = _SEV_LOOKUP.get(sev_label, "")
        fault_name = get_text(fault_key, lang)
        if canonical in ("severe",):
            recs.append(f"{fault_name}: {get_text('rec.immediate', lang)}")
        elif canonical == "moderate":
            recs.append(f"{fault_name}: {get_text('rec.maintenance', lang)}")
        elif canonical == "incipient":
            recs.append(f"{fault_name}: {get_text('rec.monitoring', lang)}")

    env_kurtosis = summary.get("envelope_kurtosis", 0)
    if env_kurtosis and env_kurtosis > 6.0:
        recs.append(f"{get_text('rec.review', lang)} ({get_text('env.kurtosis', lang)}: {env_kurtosis:.4f})")

    if not recs:
        recs.append(get_text("rec.review", lang))
    return recs


# ---------------------------------------------------------------------------
# Table helper
# ---------------------------------------------------------------------------

def _add_param_table(doc: Any, rows: list[tuple[str, str]], lang: Language) -> None:
    """Add a two-column Parameter/Value table to the document."""
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = get_text("doc.parameter", lang)
    hdr[1].text = get_text("doc.value", lang)
    for label, val in rows:
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(val)


# ---------------------------------------------------------------------------
# Section builders
# ---------------------------------------------------------------------------

def _add_motor_section(doc: Any, params: dict[str, Any], lang: Language) -> None:
    doc.add_heading(get_text("report.motor_parameters", lang), level=1)
    _add_param_table(doc, [
        (get_text("motor_param.supply_freq", lang), f'{params.get("supply_freq_hz", "—")} Hz'),
        (get_text("motor_param.poles", lang), params.get("poles", "—")),
        (get_text("motor_param.sync_speed", lang), f'{params.get("sync_speed_rpm", "—")} RPM'),
        (get_text("motor_param.rotor_speed", lang), f'{params.get("rotor_speed_rpm", "—")} RPM'),
        (get_text("motor_param.slip", lang), f'{params.get("slip", 0):.4%}' if isinstance(params.get("slip"), (int, float)) else str(params.get("slip", "—"))),
        (get_text("motor_param.rotor_freq", lang), f'{params.get("rotor_freq_hz", "—")} Hz'),
        (get_text("motor_param.slip_freq", lang), f'{params.get("slip_freq_hz", "—")} Hz'),
    ], lang)


def _add_signal_section(doc: Any, sig: dict[str, Any], source_file: str, lang: Language) -> None:
    doc.add_heading(get_text("report.signal_info", lang), level=1)
    rows = [
        (get_text("signal.n_samples", lang), f'{int(sig.get("n_samples", 0)):,}'),
        (get_text("signal.sampling_freq_hz", lang), f'{float(sig.get("sampling_freq_hz", 0)):,.0f} Hz'),
        (get_text("signal.duration_s", lang), f'{float(sig.get("duration_s", 0)):.3f} s'),
        (get_text("signal.freq_resolution_hz", lang), f'{float(sig.get("freq_resolution_hz", 0)):.6g} Hz'),
    ]
    if source_file:
        rows.append((get_text("signal.source_file", lang), source_file))
    _add_param_table(doc, rows, lang)


def _add_brb_section(doc: Any, fdata: dict[str, Any], lang: Language) -> None:
    """Broken rotor bars section."""
    from docx.shared import RGBColor

    severity = fdata.get("severity", "")
    color = _severity_color(severity)
    h = doc.add_heading(f'{get_text("fault_type.broken_rotor_bars", lang)} — {severity}', level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

    _add_param_table(doc, [
        (get_text("detail.combined_index_db", lang), f'{fdata.get("combined_index_db", "N/A")} dB'),
        (get_text("detail.severity", lang), severity),
        (get_text("detail.detected", lang), get_text(f"report.{'yes' if fdata.get('detection_status', {}).get('detected') else 'no'}", lang)),
    ], lang)

    thresholds = fdata.get("thresholds_db", {})
    if thresholds:
        doc.add_heading(get_text("detail.thresholds_db", lang), level=2)
        for sev_name in ("incipient", "moderate", "severe"):
            val = thresholds.get(sev_name)
            if val is not None:
                doc.add_paragraph(f"{sev_name}: {val} dB", style="List Bullet")


def _add_sideband_section(doc: Any, fdata: dict[str, Any], fault_label: str, lang: Language) -> None:
    """Sideband-based fault section (eccentricity / stator / bearing)."""
    from docx.shared import RGBColor

    severity = fdata.get("severity", "")
    color = _severity_color(severity)
    h = doc.add_heading(f'{fault_label} — {severity}', level=2)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

    detail_rows = [
        (get_text("detail.worst_sideband_db", lang), f'{fdata.get("worst_sideband_db", "N/A")} dB'),
        (get_text("detail.severity", lang), severity),
        (get_text("detail.detected", lang), get_text(f"report.{'yes' if fdata.get('detection_status', {}).get('detected') else 'no'}", lang)),
    ]
    if "defect_frequency_hz" in fdata:
        detail_rows.append((get_text("detail.defect_frequency_hz", lang), f'{fdata["defect_frequency_hz"]} Hz'))
    _add_param_table(doc, detail_rows, lang)

    sidebands = fdata.get("sidebands", [])
    if sidebands:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        shdr = table.rows[0].cells
        shdr[0].text = get_text("detail.harmonic_order", lang)
        shdr[1].text = get_text("detail.lower_sideband", lang)
        shdr[2].text = get_text("detail.upper_sideband", lang)
        for sb in sidebands:
            order = sb.get("harmonic_order") or sb.get("order", "")
            lo = sb.get("lower", {})
            hi = sb.get("upper", {})
            row = table.add_row().cells
            row[0].text = str(order)
            row[1].text = f'{lo.get("expected_hz", "—")} Hz'
            row[2].text = f'{hi.get("expected_hz", "—")} Hz'


def _add_env_section(doc: Any, stats: dict[str, Any], lang: Language) -> None:
    if not stats:
        return
    doc.add_heading(get_text("report.envelope_statistics", lang), level=1)
    _add_param_table(doc, [
        (get_text("env.rms", lang), f'{stats.get("rms", "—"):.6g}' if isinstance(stats.get("rms"), (int, float)) else str(stats.get("rms", "—"))),
        (get_text("env.kurtosis", lang), f'{stats.get("kurtosis", "—"):.4f}' if isinstance(stats.get("kurtosis"), (int, float)) else str(stats.get("kurtosis", "—"))),
        (get_text("env.skewness", lang), f'{stats.get("skewness", "—"):.4f}' if isinstance(stats.get("skewness"), (int, float)) else str(stats.get("skewness", "—"))),
        (get_text("env.crest_factor", lang), f'{stats.get("crest_factor", "—"):.4f}' if isinstance(stats.get("crest_factor"), (int, float)) else str(stats.get("crest_factor", "—"))),
    ], lang)


def _add_band_section(doc: Any, be: dict[str, Any], lang: Language) -> None:
    if not be:
        return
    doc.add_heading(get_text("report.band_energy", lang), level=1)
    _add_param_table(doc, [
        (get_text("band.centre_freq_hz", lang), f'{be.get("centre_freq_hz", "—")} Hz'),
        (get_text("band.bandwidth_hz", lang), f'{be.get("bandwidth_hz", be.get("band_high_hz", 0) - be.get("band_low_hz", 0))} Hz'),
        (get_text("band.band_energy", lang), f'{be.get("band_energy", "—"):.6g}' if isinstance(be.get("band_energy"), (int, float)) else str(be.get("band_energy", "—"))),
    ], lang)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def get_assessment_key_text(key: str, lang: Language) -> str:
    """Get translated assessment text from a known assessment key."""
    return get_assessment_text(key.split(".")[-1], lang)


def generate_docx_report(data: dict[str, Any], lang: Language = "en") -> tuple[bytes | None, str]:
    """Generate a DOCX diagnostic report from MCSA analysis results.

    Args:
        data: The diagnostic report dict produced by ``run_full_diagnosis`` /
            ``diagnose_from_file`` (same structure consumed by the HTML report).
        lang: Report language ("en" or "zh").

    Returns:
        Tuple of (docx_bytes, error_message).  On success error_message is
        empty; on failure docx_bytes is None.
    """
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return None, "python-docx is not installed. Install with: pip install python-docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    # Title
    doc.add_heading(get_text("report.title.mcsa_diagnostic", lang), level=0)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    doc.add_paragraph(f'{get_text("ui.generated", lang)}: {timestamp}')

    source_file = data.get("source_file", "")

    # Sections
    mp = data.get("motor_parameters", {})
    if mp:
        _add_motor_section(doc, mp, lang)

    sig = data.get("signal_info", {})
    if sig:
        _add_signal_section(doc, sig, source_file, lang)

    # Fault analysis
    fa = data.get("fault_analysis", {})
    if fa:
        doc.add_heading(get_text("report.fault_analysis", lang), level=1)

        brb = fa.get("broken_rotor_bars")
        if brb is not None:
            _add_brb_section(doc, brb, lang)
        else:
            doc.add_paragraph(f'{get_text("fault_type.broken_rotor_bars", lang)}: {get_text("report.not_computed", lang)}')

        ecc = fa.get("eccentricity")
        if ecc is not None:
            _add_sideband_section(doc, ecc, get_text("fault_type.eccentricity", lang), lang)
        else:
            doc.add_paragraph(f'{get_text("fault_type.eccentricity", lang)}: {get_text("report.not_computed", lang)}')

        stator = fa.get("stator_inter_turn")
        if stator is not None:
            _add_sideband_section(doc, stator, get_text("fault_type.stator_inter_turn", lang), lang)
        else:
            doc.add_paragraph(f'{get_text("fault_type.stator_inter_turn", lang)}: {get_text("report.not_computed", lang)}')

        bearing = fa.get("bearing")
        if bearing is not None:
            _add_sideband_section(doc, bearing, get_text("fault_type.bearing", lang), lang)

    # Band energy
    be = data.get("band_energy_around_fundamental")
    if be:
        _add_band_section(doc, be, lang)

    # Envelope statistics
    env = data.get("envelope_statistics", {})
    if env:
        _add_env_section(doc, env, lang)

    # Overall assessment
    summary = data.get("summary", {})
    overall = str(summary.get("overall_assessment", "") or "")
    if overall:
        doc.add_heading(get_text("report.overall_assessment", lang), level=1)
        # If the text is in one language but the report is in the other,
        # translate it via the key.
        key = _find_assessment_key(overall)
        if key:
            overall = get_assessment_key_text(key, lang)
        doc.add_paragraph(overall)

    # Recommendations
    doc.add_heading(get_text("report.recommendations", lang), level=1)
    for rec in _get_recommendations(summary, lang):
        doc.add_paragraph(rec, style="List Bullet")

    # Footer
    doc.add_paragraph("")
    doc.add_paragraph(f'{get_text("report.generated_by", lang)} {get_text("report.server_name", lang)}')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read(), ""
